"""Manage a local Excel workbook for the fundraising pipeline."""
import os
import re
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from config.settings import (
    EXCEL_FILE_PATH,
    EXCEL_COMPANIES_SHEET,
    EXCEL_EMAIL_LOG_SHEET,
    MANDATE_STATUSES,
)

STATUS_COLORS = {
    "Initial Contact":  "FFF9C4",
    "In Discussion":    "BBDEFB",
    "Proposal Sent":    "C8E6C9",
    "Due Diligence":    "E1BEE7",
    "Mandate Received": "A5D6A7",
    "Follow Up":        "FFE0B2",
    "On Hold":          "F5F5F5",
    "Not Interested":   "FFCDD2",
    "Closed – Won":     "1B5E20",
    "Closed – Lost":    "B71C1C",
}

_HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_THIN        = Side(style="thin", color="FFFFFF")
_BORDER      = Border(left=_THIN, right=_THIN)
_EMAIL_RE    = re.compile(r"[a-z0-9._%+-]+@([a-z0-9.-]+\.[a-z]{2,})", re.I)
_DOMAIN_RE   = re.compile(r"^(?:[a-z0-9-]+\.)+[a-z]{2,}$", re.I)
_GENERIC_COMPANY_WORDS = {
    "asset", "assets", "management", "capital", "partners", "partner",
    "group", "holdings", "holding", "limited", "ltd", "llc", "inc", "plc",
    "fund", "funds", "ventures", "venture", "family", "office",
    "investment", "investments",
}


def _write_headers(ws, headers):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _BORDER
    ws.row_dimensions[1].height = 20


def _init_companies_sheet(ws):
    _write_headers(ws, [
        "Company", "Contact Name", "Contact Email",
        "Status", "Last Email Date", "Mandate Type",
        "AUM (M€)", "Notes", "First Contact Date", "Updated",
    ])
    widths = [28, 22, 30, 20, 18, 22, 12, 40, 18, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w


def _init_email_log_sheet(ws):
    _write_headers(ws, [
        "Date", "Company", "Contact", "Direction",
        "Subject", "Preview", "Message-ID",
    ])
    widths = [18, 25, 30, 10, 45, 60, 36]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w


def _load_or_create():
    path = os.path.abspath(EXCEL_FILE_PATH)
    if os.path.exists(path):
        return openpyxl.load_workbook(path)
    wb = openpyxl.Workbook()
    ws_co = wb.active
    ws_co.title = EXCEL_COMPANIES_SHEET
    _init_companies_sheet(ws_co)
    ws_log = wb.create_sheet(EXCEL_EMAIL_LOG_SHEET)
    _init_email_log_sheet(ws_log)
    wb.save(path)
    return wb


def _save(wb):
    wb.save(os.path.abspath(EXCEL_FILE_PATH))


def _is_blank(value):
    return value is None or str(value).strip() == ""


def _normalize_key(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _company_tokens(value):
    tokens = [
        token for token in re.sub(r"[^a-z0-9]", " ", str(value or "").lower()).split()
        if len(token) > 2 and token not in _GENERIC_COMPANY_WORDS
    ]
    expanded = set(tokens)
    for token in tokens:
        if token.endswith("ical") and len(token) > 6:
            expanded.add(token[:-4])
        if token.endswith("medical") and len(token) > 8:
            expanded.add(token[:-7] + "med")
    return expanded


def _name_matches_domain(name, domain):
    domain_base = _normalize_key(str(domain or "").split(".")[0])
    if not domain_base:
        return False
    name_key = _normalize_key(name)
    if name_key == domain_base:
        return True
    return any(
        token and len(token) > 4 and (domain_base.startswith(token) or token.startswith(domain_base))
        for token in _company_tokens(name)
    )


def _email_domain(value):
    """Return an email/domain value's normalized domain, if one is present."""
    text = str(value or "").strip().lower()
    match = _EMAIL_RE.search(text)
    if match:
        return match.group(1).strip(".").lower()
    if _DOMAIN_RE.match(text):
        return text.strip(".").lower()
    return ""


def _row_to_dict(headers, row):
    return {h: row[i].value for i, h in enumerate(headers) if h}


def _company_name_score(name, domain=""):
    name = str(name or "").strip()
    if not name:
        return 0
    compact = _normalize_key(name)
    domain_base = _normalize_key((domain or "").split(".")[0])
    score = len(compact)
    if " " in name:
        score += 25
    if domain_base and compact == domain_base:
        score -= 20
    return score


def _best_company_name(existing, incoming, domain=""):
    if _is_blank(existing):
        return incoming
    if _is_blank(incoming):
        return existing
    if _company_name_score(incoming, domain) > _company_name_score(existing, domain):
        return incoming
    return existing


def _date_value(value):
    if isinstance(value, datetime):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    text = text.replace("T", " ").replace("Z", "")
    for fmt, length in (
        ("%Y-%m-%d %H:%M:%S", 19),
        ("%Y-%m-%d %H:%M", 16),
        ("%Y-%m-%d", 10),
    ):
        try:
            return datetime.strptime(text[:length], fmt)
        except ValueError:
            continue
    return None


def _merge_domain_match_data(data, existing, headers, domain):
    """Merge an incoming auto-classification with an existing same-domain row.

    Domain matches often mean the model inferred a shorter name from the email
    domain, while the workbook may already contain a manually corrected name.
    Preserve richer existing values when the incoming payload is blank, and keep
    the better display name instead of creating a second deal.
    """
    merged = dict(data)
    merged["Company"] = _best_company_name(existing.get("Company"), data.get("Company"), domain)

    for h in headers:
        if h not in merged or h == "Company":
            continue
        incoming = merged.get(h)
        current = existing.get(h)
        if _is_blank(incoming) and not _is_blank(current):
            merged.pop(h, None)
        elif h == "Last Email Date":
            incoming_date = _date_value(incoming)
            current_date = _date_value(current)
            if current_date and (not incoming_date or current_date > incoming_date):
                merged[h] = current
    return merged


def _apply_row_color(ws, row_num, col_count, status):
    color = STATUS_COLORS.get(status, "FFFFFF")
    fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    dark = color in ("1B5E20", "B71C1C")
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = fill
        cell.font = Font(color="FFFFFF" if dark else "000000")


# ── Public API ─────────────────────────────────────────────────────────────

def get_companies():
    wb = _load_or_create()
    ws = wb[EXCEL_COMPANIES_SHEET]
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            rows.append(dict(zip(headers, row)))
    return rows


def upsert_company(data, match_name=None):
    """Insert or update a company row.

    Existing rows are matched by ``match_name`` when supplied, by the incoming
    Company value, or by the contact/domain carried by the email thread. This
    lets the UI rename a deal without leaving the old company row behind, and
    prevents domain-derived aliases from creating duplicate deals.
    """
    data = dict(data)
    wb = _load_or_create()
    ws = wb[EXCEL_COMPANIES_SHEET]
    headers = [c.value for c in ws[1]]
    data["Updated"] = datetime.now().strftime("%Y-%m-%d %H:%M")

    target = None
    lookup_names = [
        str(name).strip().lower()
        for name in (match_name, data.get("Company", ""))
        if str(name or "").strip()
    ]
    incoming_domain = (
        _email_domain(data.get("_Domain")) or
        _email_domain(data.get("Domain")) or
        _email_domain(data.get("Contact Email"))
    )

    for row in ws.iter_rows(min_row=2):
        row_name = str(row[0].value or "").strip().lower()
        if row_name and row_name in lookup_names:
            target = row[0].row
            break

    if not target and incoming_domain:
        email_idx = headers.index("Contact Email") if "Contact Email" in headers else None
        for row in ws.iter_rows(min_row=2):
            row_domain = _email_domain(row[email_idx].value) if email_idx is not None else ""
            if (row_domain and row_domain == incoming_domain) or _name_matches_domain(row[0].value, incoming_domain):
                target = row[0].row
                existing = _row_to_dict(headers, row)
                data = _merge_domain_match_data(data, existing, headers, incoming_domain)
                break

    if target:
        for col, h in enumerate(headers, 1):
            if h in data:
                ws.cell(row=target, column=col, value=data[h])
    else:
        if "First Contact Date" not in data:
            data["First Contact Date"] = datetime.now().strftime("%Y-%m-%d")
        ws.append([data.get(h, "") for h in headers])
        target = ws.max_row

    status_col = headers.index("Status") + 1 if "Status" in headers else 0
    status = ws.cell(row=target, column=status_col).value if status_col else data.get("Status", "")
    _apply_row_color(ws, target, len(headers), status or "")
    _save(wb)


def update_company_status(company_name, new_status, notes=None):
    patch = {"Company": company_name, "Status": new_status}
    if notes:
        patch["Notes"] = notes
    upsert_company(patch)


def log_emails(emails, company_name):
    """Append emails to Email Log sheet, skipping duplicates by Message-ID."""
    wb = _load_or_create()
    ws = wb[EXCEL_EMAIL_LOG_SHEET]

    existing_ids = {
        str(row[6]) for row in ws.iter_rows(min_row=2, values_only=True) if row[6]
    }

    added = 0
    for em in emails:
        msg_id = str(em.get("id") or "")
        if msg_id and msg_id in existing_ids:
            continue
        contact = (em.get("external_participants") or [em.get("from") or {}])[0]
        ea = (contact or {}).get("emailAddress", {})
        contact_addr = ea.get("address", "")
        direction = em.get("direction") or ("OUT" if (em.get("folder") or "").lower() == "sent" else "IN")
        ws.append([
            (em.get("receivedDateTime") or "")[:10],
            company_name,
            f"{ea.get('name', '')} <{contact_addr}>".strip(),
            direction,
            em.get("subject", ""),
            (em.get("bodyPreview") or "")[:200],
            msg_id,
        ])
        existing_ids.add(msg_id)
        added += 1

    if added:
        _save(wb)
    return added


def get_excel_path():
    return os.path.abspath(EXCEL_FILE_PATH)
